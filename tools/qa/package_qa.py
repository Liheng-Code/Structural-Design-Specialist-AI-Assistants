#!/usr/bin/env python3
"""Cross-document QA of a structural document package. Produces an issue register; it never approves anything.

Usage:  python tools/qa/package_qa.py <package_dir> [--out <dir>] [--standards <Structural-Drawing-Standards>] [--db db.yaml] [--register reg.csv]

Package layout (convention):  <dir>/design_basis.yaml  <dir>/document_register.csv  <dir>/output/**  (DOCX, PDF, XLSX, DXF, CALC-*.json)
                              <dir>/analysis/*/import_log.json  (optional, from tools/etabs/etabs_import.py)
Document files are named <NUMBER>_<REV>.<ext>, e.g. WH-STR-CR-001_R00.docx.

Writes issue-register.csv (columns of the QA skill template) and qa_report.md. Exit 0 = no open Critical/Major finding
(the CONSISTENCY gate). That is not issue approval: issue readiness is reported separately and always needs authorised
engineering review. Everything the tool could not check is listed in the report.

Checks: Q1 design basis | Q2 register | Q3 files vs register | Q4 identity on DOCX/XLSX/DXF/PDF | Q5 status claims and DRAFT stamps |
Q6 placeholders | Q7 report numbers vs results JSON vs workbook | Q8 inputs vs design-basis materials | Q9 DXF standards validation |
Q10 workbook integrity | Q11 file integrity | Q12 ETABS/SAFE import | Q13 open items / editions / sign-off.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import pathlib
import re
import subprocess
import sys

import yaml

HERE = pathlib.Path(__file__).resolve().parent
TOOLS = HERE.parent
for sub in ("design_basis", "doc_control", "xlsxgen", "cadgen"):
    sys.path.insert(0, str(TOOLS / sub))
import check_register as cr  # noqa: E402
import validate_design_basis as vdb  # noqa: E402

FILE_RE = re.compile(r"^(?P<num>[A-Z]{2,4}-STR-[A-Z]{2}-\d{3})_(?P<rev>[RPC]\d{2})(?:_preview)?\.(?P<ext>docx|pdf|xlsx|dxf)$")
PLACE_RE = re.compile(r"\b(TBC|TBD|XXX+)\b", re.I)   # tokens only; the words 'placeholder' / 'to be confirmed' occur legitimately in open-item text
CRIT, MAJ, MIN, INFO = "Critical", "Major", "Minor", "Info"
COLS = ["issue_id", "severity", "document", "location", "finding", "evidence", "corrective_action", "owner", "review_status", "closed_date"]


class QA:
    def __init__(self):
        self.f, self.not_checked = [], []

    def add(self, sev, doc, loc, finding, evidence, action, check):
        self.f.append({"severity": sev, "document": doc, "location": loc, "finding": f"[{check}] {finding}", "evidence": evidence,
                       "corrective_action": action, "owner": "", "review_status": "Open", "closed_date": ""})

    def nc(self, check, why):
        self.not_checked.append(f"{check}: {why}")


def docx_text(path):
    from docx import Document
    d = Document(path)
    paras = [p.text for p in d.paragraphs]
    tables = [[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
    hdr = " ".join(p.text for s in d.sections for p in s.header.paragraphs)
    return d, paras, tables, hdr


def ctrl_map(tables):
    for t in tables:
        if t and t[0][:2] == ["Item", "Value"]:
            return {r[0]: r[1] for r in t[1:]}
    return {}


def jload(p):
    with open(p, encoding="utf-8") as fh:
        return json.load(fh)


def sha(p):
    return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()


def run(pkg, standards=None, db_path=None, reg_path=None) -> QA:
    q = QA()
    pkg = pathlib.Path(pkg)
    dbp = pathlib.Path(db_path) if db_path else pkg / "design_basis.yaml"
    regp = pathlib.Path(reg_path) if reg_path else pkg / "document_register.csv"
    if not dbp.exists():
        q.add(CRIT, "-", str(dbp), "design basis file is missing", "file not found", "Provide design_basis.yaml; no consistency check is possible without it", "Q1")
        return q
    with open(dbp, encoding="utf-8") as fh:
        db = yaml.safe_load(fh)
    P = db.get("project", {})
    code, name, client, stage, rev, date, dbstatus = (P.get(k) for k in ("code", "name", "client", "stage", "revision", "issue_date", "status"))
    dbstatus = (db.get("approval") or {}).get("status")

    # Q1 design basis
    for sev, path, msg in vdb.validate(db):
        if sev != "ERROR" and re.search(r"edition/year not recorded|generation \(1st/2nd\) not stated|open item|unresolved placeholder", msg):
            continue   # summarised once under Q13 / Q6
        q.add(MAJ if sev == "ERROR" else INFO, "design_basis.yaml", path, msg, f"validate_design_basis: {sev}", "Resolve in the design basis" if sev == "ERROR" else "Review; may be an accepted open item", "Q1")

    # Q2 register
    reg = None
    if regp.exists():
        with open(regp, newline="", encoding="utf-8-sig") as fh:
            reg = list(csv.DictReader(fh))
        for sev, ref, msg in cr.check(reg, db):
            q.add(MAJ if sev == "ERROR" else MIN, "document_register.csv", ref, msg, f"check_register {sev}", "Correct the register or the sheet", "Q2")
    else:
        q.add(MAJ, "-", str(regp), "document register is missing", "file not found", "Provide the register", "Q2")

    out = pkg / "output"
    files = sorted(p for p in out.rglob("*") if p.is_file()) if out.exists() else []
    docs = {}
    for p in files:
        m = FILE_RE.match(p.name)
        if m:
            docs.setdefault((m["num"], m["rev"]), {})[m["ext"] + ("-preview" if "_preview" in p.name else "")] = p
        elif p.suffix.lower() in (".docx", ".xlsx", ".dxf") :
            q.add(MIN, p.name, "file name", "file name does not follow <NUMBER>_<REV>.<ext>", p.name, "Rename per document control", "Q3")
    if not files:
        q.add(MAJ, "-", str(out), "no output files in package", "empty", "Nothing to review", "Q3")

    # Q3 files vs register / numbering
    reg_nums = {r["sheet_number"]: r for r in reg} if reg else {}
    for (num, r), d in docs.items():
        if db.get("project", {}).get("drawing_prefix") and not num.startswith(P["drawing_prefix"] + "-"):
            q.add(MAJ, f"{num}_{r}", "number", f"document number prefix differs from design basis drawing prefix {P['drawing_prefix']}", num, "Renumber or correct the design basis", "Q3")
        if "-SD-" in num:
            if reg and num not in reg_nums:
                q.add(MAJ, f"{num}_{r}", "register", "drawing is not in the document register", num, "Add to register or remove the file", "Q3")
            elif reg and reg_nums[num]["revision"] != r:
                q.add(MAJ, f"{num}_{r}", "register", f"file revision {r} != register revision {reg_nums[num]['revision']}", num, "Align revision", "Q3")
        elif reg is not None:
            q.add(INFO, f"{num}_{r}", "register", "non-drawing document is not tracked in the drawing register", num, "Track in a document list if the project requires it", "Q3")
    if reg:
        have = {n for n, _ in docs}
        missing = sorted(set(reg_nums) - have)
        if missing:
            q.add(INFO, "document_register.csv", "sheets without files", f"{len(missing)} register sheet(s) have no file in this package", ", ".join(missing[:6]) + (" ..." if len(missing) > 6 else ""),
                  "Not an error if the sheets are issued elsewhere; they were NOT checked", "Q3")
            q.nc("Q3-Q9", f"{len(missing)} register sheets have no file in the package")

    doc_blob = {}  # number -> text, for Q6
    # Q4/Q5/Q11 per document
    for (num, r), d in sorted(docs.items()):
        tag = f"{num}_{r}"
        if "docx" in d:
            try:
                _, paras, tables, hdr = docx_text(d["docx"])
            except Exception as ex:  # noqa: BLE001
                q.add(CRIT, tag, "docx", "DOCX cannot be opened", str(ex), "Regenerate", "Q11")
                continue
            c = ctrl_map(tables)
            doc_blob[tag] = "\n".join(paras + [x for t in tables for rr in t for x in rr] + [hdr])
            if not c:
                q.add(MAJ, tag, "document control table", "no document control table found", "first table is not Item/Value", "Add the control table", "Q4")
            else:
                exp = {"Document number": num, "Project": f"{code} {name}", "Client": client, "Stage": stage, "Revision / date": f"{rev} / {date}"}
                for k, v in exp.items():
                    if c.get(k) != v:
                        q.add(MAJ, tag, f"control table: {k}", f"'{c.get(k)}' != expected '{v}'", f"{tag}.docx vs design basis/file name", "Regenerate from the current design basis", "Q4")
                if r != rev:
                    q.add(MIN, tag, "revision", f"file revision {r} differs from design basis revision {rev}", r, "OK only if this document was individually revised", "Q4")
                for k in ("Checked by", "Approved by"):
                    if c.get(k, "").strip():
                        q.add(INFO, tag, k, f"'{k}' is filled in ('{c[k]}')", "QA cannot verify this is a real human review", "Confirm with the named person", "Q13")
                if "Design basis" in c and dbstatus and dbstatus not in c["Design basis"]:
                    q.add(MAJ, tag, "control table: Design basis", f"states '{c['Design basis']}' but design basis status is {dbstatus}", "", "Regenerate", "Q4")
            if code and code not in hdr:
                q.add(MAJ, tag, "header", "page header does not carry the project code", hdr, "Regenerate", "Q4")
            if dbstatus == "DRAFT" and "DRAFT" not in hdr.upper():
                q.add(MAJ, tag, "header", "design basis is DRAFT but the document is not stamped DRAFT", hdr, "Stamp DRAFT - NOT FOR ISSUE", "Q5")
            if re.search(r"\b(APPROVED FOR (ISSUE|CONSTRUCTION)|VALIDATED)\b", doc_blob[tag], re.I) and dbstatus != "APPROVED":
                q.add(CRIT, tag, "text", "document claims approval/validation while the design basis is not APPROVED", "approval wording found", "Remove the claim; only a named engineer can approve", "Q5")
        if "xlsx" in d:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(d["xlsx"], data_only=True)
                cv = {row[0]: row[1] for row in wb["Cover"].iter_rows(values_only=True) if row and row[0]}
            except Exception as ex:  # noqa: BLE001
                q.add(CRIT, tag, "xlsx", "workbook cannot be opened or has no Cover sheet", str(ex), "Regenerate", "Q11")
                cv = {}
            if cv:
                exp = {"Workbook number": num, "Project": f"{code} {name}", "Client": client, "Stage": stage, "Revision": rev, "Issue date": date, "Design basis status": dbstatus}
                for k, v in exp.items():
                    if cv.get(k) != v:
                        q.add(MAJ, tag, f"Cover: {k}", f"'{cv.get(k)}' != expected '{v}'", "", "Regenerate", "Q4")
                st = str(cv.get("STATUS", ""))
                if re.search(r"VALIDATED|APPROVED|REVIEWED|ISSUED", st) and not re.search(r"NOT", st):
                    q.add(CRIT, tag, "Cover: STATUS", f"workbook status '{st}' claims a review/approval", st, "A status needs a named reviewer", "Q5")
                if dbstatus == "DRAFT" and "DRAFT" not in str(cv.get("Stamp", "")).upper():
                    q.add(MAJ, tag, "Cover: Stamp", "design basis DRAFT but workbook is not stamped DRAFT", str(cv.get("Stamp")), "Stamp", "Q5")
        if "pdf" in d:
            try:
                txt = subprocess.run(["pdftotext", "-l", "1", str(d["pdf"]), "-"], capture_output=True, text=True, timeout=60, check=True).stdout
                if num not in txt:
                    q.add(MAJ, tag, "pdf page 1", "PDF does not show the document number", "", "Regenerate the PDF from the DOCX", "Q11")
                if "docx" in d and not txt.strip():
                    q.add(MAJ, tag, "pdf", "PDF has no text", "", "Regenerate", "Q11")
            except (OSError, subprocess.SubprocessError):
                q.nc("Q11", f"pdftotext not available: PDF content of {tag} not checked")
        elif "docx" in d:
            q.add(MIN, tag, "pdf", "no PDF companion for the DOCX", "", "Generate the PDF", "Q11")
        if "dxf" in d:
            import validate_dxf as vd
            std = pathlib.Path(standards) if standards else None
            if std and std.exists():
                res = vd.validate(d["dxf"], std, db, reg)
                for e in res["errors"]:
                    q.add(MAJ, tag, "dxf", e, "validate_dxf", "Fix the generator or the sheet content", "Q9")
                tb = res.get("title_block", {})
                if tb.get("REVISION") and tb["REVISION"] != r:
                    q.add(MAJ, tag, "title block REVISION", f"{tb['REVISION']} != file revision {r}", "", "Align", "Q4")
                if dbstatus == "DRAFT" and "DRAFT" not in str(tb).upper():
                    q.add(INFO, tag, "title block", "design basis DRAFT; sheet carries no DRAFT marking", "", "Mark the sheet or confirm that the register status governs", "Q5")
            else:
                q.nc("Q9", "no --standards folder: DXF layer/text/title-block rules not validated")
            q.add(INFO, tag, "format", "CAD deliverable is DXF only (no DWG); preview PDF is not an AutoCAD plot", d["dxf"].name, "State this on transmittals", "Q9")

    # Q6 placeholders
    inherited = {m.group(0).upper() for v in (client, name, P.get("location")) if isinstance(v, str) for m in PLACE_RE.finditer(v)}
    for k in ("client", "name", "location", "engineer"):
        v = P.get(k)
        if isinstance(v, str) and PLACE_RE.search(v):
            q.add(MIN if dbstatus == "DRAFT" else MAJ, "design_basis.yaml", f"project.{k}", f"placeholder in '{v}'", v, "Replace before issue", "Q6")
    for tag, txt in doc_blob.items():
        for m in sorted({m.group(0).upper() for m in PLACE_RE.finditer(txt)}):
            if m not in inherited:
                q.add(MAJ, tag, "text", f"placeholder '{m}' not traceable to the design basis", m, "Resolve or add to the open items", "Q6")
            else:
                q.add(INFO, tag, "text", f"placeholder '{m}' inherited from the design basis", m, "Resolve with the design basis", "Q6")

    # Q7/Q8 results JSON vs report vs workbook; inputs vs materials
    mats = {m["id"]: m for m in db.get("materials", []) if isinstance(m, dict) and "id" in m}
    for jp in sorted(out.glob("CALC-*.json")) if out.exists() else []:
        try:
            res = jload(jp)
        except Exception as ex:  # noqa: BLE001
            q.add(CRIT, jp.name, "json", "results file unreadable", str(ex), "Re-run the calculation", "Q7")
            continue
        cid = res.get("id", jp.stem)
        if not res.get("verification"):
            q.add(CRIT, cid, "verification", "result has no verification record (PASS/values without independent basis)", "", "Add hand/independent check", "Q7")
        # Q8
        for inp in res.get("inputs", []):
            m = re.search(r"design basis (M-[A-Z0-9-]+)", str(inp.get("source", "")))
            if not m:
                continue
            mat = mats.get(m[1])
            if mat is None:
                q.add(MAJ, cid, f"input {inp['symbol']}", f"cites design-basis item {m[1]} which does not exist", inp.get("source"), "Fix the reference", "Q8")
                continue
            spec = str(mat.get("spec", ""))
            sym = inp["symbol"]
            if sym == "concrete class":
                if str(inp["value"]) not in spec:
                    q.add(MAJ, cid, f"input {sym}", f"{inp['value']} not in design basis {m[1]} ('{spec}')", "", "Align", "Q8")
            elif sym == "fyk":
                f = re.search(r"fyk\s*=\s*(\d+)", spec)
                if not f or float(f[1]) != float(inp["value"]):
                    q.add(MAJ, cid, f"input {sym}", f"{inp['value']} != design basis {m[1]} ('{spec}')", "", "Align", "Q8")
        rows = {"lb,rqd good bond (mm)": "lb_rqd_good", "lb,rqd poor bond (mm)": "lb_rqd_poor", "lap l0 good, 50% lapped (mm)": "lap_basic_good"}
        for tag, d in docs.items():
            if "docx" not in d or cid not in doc_blob.get(f"{tag[0]}_{tag[1]}", ""):
                continue
            _, _, tables, _ = docx_text(d["docx"])
            found = 0
            for t in tables:
                for row in t:
                    key = rows.get(row[0])
                    if key:
                        found += 1
                        want = res["results"][key]
                        try:
                            got = [float(x) for x in row[1:]]
                        except ValueError:
                            got = None
                        # the report prints whole millimetres; accept either rounding of an exact .5
                        if got is None or len(got) != len(want) or any(abs(g - w) > 0.5 + 1e-9 for g, w in zip(got, want)):
                            q.add(CRIT, f"{tag[0]}_{tag[1]}", f"table row '{row[0]}'", f"report values {got} != results file {cid} {[round(w, 1) for w in want]}", "", "Regenerate the report from the results file", "Q7")
            if found < len(rows):
                q.add(MAJ, f"{tag[0]}_{tag[1]}", "results tables", f"references {cid} but only {found}/{len(rows)} result rows were found to compare", "", "Check report structure", "Q7")
        # workbook tie-out (no recalculation: cached values must be present)
        for tag, d in docs.items():
            if "xlsx" in d:
                try:
                    import verify_workbook as vw
                    v = vw.verify(d["xlsx"], res, do_recalc=False)
                    for pr in v["problems"]:
                        q.add(MIN if "data validation" in pr else CRIT, f"{tag[0]}_{tag[1]}", "workbook", pr, "verify_workbook (no recalc)", "Fix and recalculate", "Q10")
                except Exception as ex:  # noqa: BLE001
                    q.nc("Q10", f"workbook {tag[0]} could not be verified: {ex}")
    if out.exists() and not list(out.glob("CALC-*.json")) and any(n for n in docs if "-CR-" in n[0] or "-WB-" in n[0]):
        q.add(MAJ, "-", "results", "calculation report/workbook present but no CALC-*.json results file to tie out against", "", "Provide the results file", "Q7")
        q.nc("Q7", "no results JSON: numbers not tied out")

    # Q12 ETABS/SAFE
    logs = sorted((pkg / "analysis").glob("*/import_log.json")) if (pkg / "analysis").exists() else []
    if not logs:
        q.nc("Q12", "no ETABS/SAFE import in the package: analysis provenance, equilibrium and mapping not checked")
    for lp in logs:
        lg = jload(lp)
        a = lp.parent.name
        prov = lg["provenance"]
        is_db = "model_definition_usable" in lg.get("summary", {})
        files_ = prov.get("export_files") or [{"file": prov.get("file"), "sha256": prov.get("sha256")}]
        if is_db:
            if not lg["summary"]["model_definition_usable"]:
                q.add(CRIT, a, "import_log", "model .db import has failing checks", f"checks_fail={lg['summary'].get('checks_fail')}", "Resolve the failures", "Q12")
            q.add(MAJ, a, "import_log", f"import holds a model definition only ({lg['summary'].get('reason_not_ready')}); no analysis results are available to support calculations", "", "Supply a results export before relying on analysis values", "Q12")
        elif not lg.get("summary", {}).get("ready_for_calculations"):
            q.add(CRIT, a, "import_log", "import is not ready for calculations but the package contains it", f"blocking_problems={lg['summary'].get('blocking_problems')}", "Resolve import failures first", "Q12")
        for e in files_:
            rp = lp.parent / "raw" / e["file"]
            if not rp.exists():
                if is_db:
                    q.nc("Q12", f"raw copy of {e['file']} is not in the package: its SHA-256 could not be re-verified")
                else:
                    q.add(CRIT, a, f"raw/{e['file']}", "raw export is missing", "", "Restore the unchanged export", "Q12")
            elif sha(rp) != e["sha256"]:
                q.add(CRIT, a, f"raw/{e['file']}", "raw export differs from the recorded SHA-256", "", "Restore the unchanged export", "Q12")
        for c in lg["checks"]:
            if c["status"] == "NOT CHECKED":
                q.add(MAJ, a, c["id"], f"import check not performed: {c['name']}", c["detail"], "Supply the missing input and re-run", "Q12")
            elif c["status"] == "FAIL":
                q.add(CRIT, a, c["id"], c["name"], c["detail"], "Investigate the model/export", "Q12")
        ap = db.get("analysis_provenance") or {}
        miss = [k for k in ("software", "model_file", "export_sha256") if not ap.get(k)]
        if miss:
            q.add(MAJ, "design_basis.yaml", "analysis_provenance", f"analysis results are in the package but the design basis provenance lacks {miss}", a,
                  "Review analysis_provenance.proposed.yaml and record it in the design basis", "Q12")
        else:
            hs = {e["file"]: e["sha256"] for e in files_}
            if ap["export_sha256"] != hs:
                q.add(CRIT, "design_basis.yaml", "analysis_provenance", "recorded export hashes differ from the imported files", "", "Re-import or correct provenance", "Q12")

    # Q13 open items / editions / sign-off / readiness
    open_items = [o for o in db.get("open_items", []) if str(o.get("status", "open")).lower() == "open"]
    blank_ed = [s.get("id") for s in db.get("standards", []) if isinstance(s, dict) and (not s.get("edition") or not s.get("generation"))]
    if open_items:
        q.add(MAJ if dbstatus == "APPROVED" else INFO, "design_basis.yaml", "open_items", f"{len(open_items)} open item(s): " + ", ".join(o["id"] for o in open_items), "", "Close or accept (engineer)", "Q13")
    if blank_ed:
        q.add(MAJ if dbstatus == "APPROVED" else INFO, "design_basis.yaml", "standards", f"edition/generation not recorded for {len(blank_ed)} standard(s)", ", ".join(map(str, blank_ed[:6])), "Record before approval", "Q13")
    return q


def summarise(q: QA, db_status=None):
    open_blockers = [f for f in q.f if f["severity"] in (CRIT, MAJ) and f["review_status"] == "Open"]
    counts = {s: sum(f["severity"] == s for f in q.f) for s in (CRIT, MAJ, MIN, INFO)}
    return {"counts": counts, "consistency_gate_pass": not open_blockers, "not_checked": q.not_checked}


def write(q: QA, outdir, pkg, db_status):
    outdir = pathlib.Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    order = {CRIT: 0, MAJ: 1, MIN: 2, INFO: 3}
    q.f.sort(key=lambda f: order[f["severity"]])
    for i, f in enumerate(q.f, 1):
        f["issue_id"] = f"QA-{i:03d}"
    with open(outdir / "issue-register.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS)
        w.writeheader()
        w.writerows(q.f)
    s = summarise(q)
    L = [f"# QA report - {pkg}", "", f"Consistency gate (no open Critical/Major): **{'PASS' if s['consistency_gate_pass'] else 'FAIL'}**",
         f"Issue readiness: **NOT READY** - design basis status {db_status}; authorised engineering review is required before final issue. This report is not an approval.", "",
         "Findings: " + ", ".join(f"{k} {v}" for k, v in s["counts"].items()), "", "## Not checked / limits"]
    L += [f"- {x}" for x in s["not_checked"]] or ["- (none recorded)"]
    L += ["- Engineering correctness of design inputs and results is not judged; only consistency and traceability.", "", "## Findings", "| Id | Severity | Document | Location | Finding |", "|---|---|---|---|---|"]
    L += [f"| {f['issue_id']} | {f['severity']} | {f['document']} | {f['location']} | {f['finding'].replace('|', '/')} |" for f in q.f]
    (outdir / "qa_report.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    return s


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("package")
    ap.add_argument("--out")
    ap.add_argument("--standards")
    ap.add_argument("--db")
    ap.add_argument("--register")
    a = ap.parse_args(argv)
    q = run(a.package, a.standards, a.db, a.register)
    dbp = pathlib.Path(a.db) if a.db else pathlib.Path(a.package) / "design_basis.yaml"
    st = (yaml.safe_load(dbp.read_text(encoding="utf-8")) or {}).get("approval", {}).get("status") if dbp.exists() else "unknown"
    s = write(q, a.out or pathlib.Path(a.package) / "qa", a.package, st)
    print(json.dumps({"counts": s["counts"], "consistency_gate_pass": s["consistency_gate_pass"], "not_checked": len(s["not_checked"])}, indent=2))
    return 0 if s["consistency_gate_pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
