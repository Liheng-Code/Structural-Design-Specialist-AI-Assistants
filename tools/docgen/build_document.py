#!/usr/bin/env python3
"""Build a structural specification or calculation report (DOCX + PDF) from a content file and the design basis.

Usage:
  python tools/docgen/build_document.py <content.yaml> --design-basis <db.yaml> --out <dir> [--results a.json b.json]

Design rules enforced here (not left to the author):
  * Project code, name, revision, date, stage, prepared-by are taken from the design basis only.
  * Checked-by / approved-by are never filled in. They print as blank signature cells.
  * Documents built from a design basis whose status is not APPROVED are stamped "DRAFT - NOT FOR ISSUE".
  * Every specification clause must carry a source tag; a clause without one stops the build.
  * Headings/clauses are numbered by the program, so numbering is continuous by construction.
  * Open items from the design basis are always appended as 'Assumptions to be confirmed'.
Requires: python-docx, PyYAML; LibreOffice (soffice) for PDF.
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import shutil
import subprocess
import sys

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PLACEHOLDER = re.compile(r"\b(TBC|TBD|XXX+|TO BE CONFIRMED|PLACEHOLDER)\b", re.I)
FONT = "Arial"
GREY = RGBColor(0x59, 0x59, 0x59)


class BuildError(Exception):
    pass


# ----------------------------------------------------------------------------- helpers
def _font(run, size=10, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    if color is not None:
        run.font.color.rgb = color


def _para(doc, text="", size=10, bold=False, italic=False, align=None, space_after=4, color=None, keep_next=False):
    p = doc.add_paragraph()
    if text:
        _font(p.add_run(text), size, bold, italic, color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = keep_next
    if align:
        p.alignment = align
    return p


def _shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _table(doc, header, rows, widths=None, size=8.5, caption=None):
    if caption:
        _para(doc, caption, 9, bold=True, space_after=2, keep_next=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(str(h)), size, bold=True)
        _shade(c, "E7E6E6")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run("" if v is None else str(v)), size)
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    if widths is None and len(header) > 2:  # sensible default: first column wider, rest equal
        first = 58 if len(header) > 4 else 45
        widths = [first] + [(170 - first) / (len(header) - 1)] * (len(header) - 1)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Mm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Mm(w)
    _para(doc, "", space_after=4)
    return t


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _letters(n: int) -> str:
    return chr(ord("A") + n)


# ----------------------------------------------------------------------------- checks
def check_content(content: dict, db: dict) -> tuple[list, list]:
    errors, warnings = [], []
    kind = content["doc"]["type"]
    for si, sec in enumerate(content.get("sections", []), start=1):
        for bi, b in enumerate(sec.get("blocks", []), start=1):
            where = f"section {si} '{sec['title']}' block {bi}"
            if b["type"] == "clause" and kind == "specification" and not b.get("source"):
                errors.append(f"{where}: clause has no source tag: '{b['text'][:50]}'")
            for key in ("text", "source"):
                if key in b and PLACEHOLDER.search(str(b[key])):
                    warnings.append(f"{where}: placeholder in {key}: '{str(b[key])[:50]}'")
            if b["type"] == "clause" and b.get("status") == "seed":
                warnings.append(f"{where}: clause is a SEED from the issued drawings - verify before issue")
    if db["approval"]["status"] != "APPROVED":
        warnings.append("design basis is not APPROVED - document stamped DRAFT - NOT FOR ISSUE")
    return errors, warnings


# ----------------------------------------------------------------------------- block renderers
def render_calc(doc, res: dict):
    _para(doc, f"{res['id']}  {res['title']}", 10, bold=True, keep_next=True)
    _para(doc, f"Code reference: {res['clause']}", 9, italic=True, color=GREY)
    _table(doc, ["Input", "Value", "Unit", "Source"],
           [[i["symbol"], i["value"], i["unit"], i["source"]] for i in res["inputs"]],
           widths=[35, 25, 15, 95], caption="Inputs")
    _table(doc, ["Step", "Formula", "Substitution", "Result", "Unit"],
           [[s["label"], s["formula"], s["substitution"], s["result"], s["unit"]] for s in res["steps"]],
           widths=[34, 56, 36, 16, 14], caption="Calculation")
    bars = res["bars_mm"]
    hdr = ["Bar diameter (mm)"] + [f"{b:g}" for b in bars]
    r = res["results"]
    _table(doc, hdr, [["lb,rqd good bond (mm)"] + [f"{x:.0f}" for x in r["lb_rqd_good"]],
                      ["lb,rqd poor bond (mm)"] + [f"{x:.0f}" for x in r["lb_rqd_poor"]],
                      ["lap l0 good, 50% lapped (mm)"] + [f"{x:.0f}" for x in r["lap_basic_good"]]],
           caption="Results (unrounded; tabulated fctk,0.05)")
    _para(doc, f"Scope limits: {res['scope_limits']}", 8.5, italic=True)
    rec = res.get("reconciliation")
    if rec:
        rows = []
        for x in rec["rows"]:
            rows.append([x["row"] + " - drawing"] + x["drawing"])
            rows.append([x["row"] + " - calc, drawing basis (fctk,0.05 = 2.03)"] + x["calc_drawing_basis"])
            rows.append(["   deviation (drawing - calc), mm"] + [f"{d:+d}" for d in x["dev_drawing_basis"]])
            rows.append([x["row"] + " - calc, Table 3.1 (fctk,0.05 = 2.0)"] + x["calc_table_value"])
            rows.append(["   deviation (drawing - calc), mm"] + [f"{d:+d}" for d in x["dev_table_value"]])
        _table(doc, hdr, rows, caption=f"Reconciliation with {rec['label']}", size=8)
        _para(doc, rec["note"], 8.5, italic=True)
        _para(doc, "Drawing values agree with the calculation on the drawing's own basis within "
                   f"{rec['tolerance']}: {'YES' if rec['all_within_tolerance_drawing_basis'] else 'NO'}. "
                   f"Against the tabulated fctk,0.05 = 2.0 MPa the drawing values are shorter by up to "
                   f"{rec['max_rel_dev_table_value_pct']} % ({rec['max_dev_table_value_mm']} mm).", 9)
        m = rec.get("mandrel")
        if m and m["mismatch_bars"]:
            _para(doc, "FINDING - mandrel diameters: the drawing's hook/bend row gives "
                       f"{m['drawing']} mm for bars {[int(b) for b in res['bars_mm']]} mm; EN 1992-1-1 8.3(1) Table 8.1 "
                       f"(4 phi for phi <= 16 mm, 7 phi above) gives {m['en_1992_8_3']} mm. "
                       f"Mismatch for bars {m['mismatch_bars']} mm diameter. The same drawing's note and sheet SD-006 state 7 phi "
                       "for phi > 16 mm. Engineer to confirm and correct SD-004 at next revision.", 9, bold=True)
    _para(doc, f"Verification status: {res['verification']['status']}. Method: {res['verification']['method']}.", 8.5, italic=True)


def render_block(doc, b, results, numbering):
    t = b["type"]
    if t == "paragraph":
        _para(doc, b["text"], 10)
    elif t == "note":
        _para(doc, "Note: " + b["text"], 9, italic=True, color=GREY)
    elif t == "clause":
        numbering["clause"] += 1
        label = f"{numbering['section']}.{numbering['clause']}" if numbering["style"] == "decimal" else f"{numbering['clause']}."
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(12)
        p.paragraph_format.first_line_indent = Mm(-12)
        p.paragraph_format.space_after = Pt(4)
        _font(p.add_run(label + "\t"), 10, bold=True)
        _font(p.add_run(b["text"] + " "), 10)
        _font(p.add_run(f"[{b['source']}]"), 8.5, italic=True, color=GREY)
        if b.get("status") == "seed":
            _font(p.add_run("  [SEED - verify]"), 8.5, bold=True, color=RGBColor(0xC0, 0x00, 0x00))
    elif t == "table":
        _table(doc, b["columns"], b["rows"], caption=b.get("caption"), widths=b.get("widths"))
    elif t == "db_standards":
        db = numbering["db"]
        _table(doc, ["Standard", "Generation", "Edition", "Used for"],
               [[x["id"], x.get("generation") or "NOT RECORDED", x.get("edition") or "NOT RECORDED", x.get("used_for", "")]
                for x in db["standards"]], widths=[30, 28, 28, 84], caption="Design codes (from design basis)")
        if db.get("mixed_generation_justification"):
            _para(doc, "Mixed 1st/2nd generation Eurocodes: " + db["mixed_generation_justification"], 9, italic=True)
        na = db["national_annex"]
        _para(doc, f"National Annex: {na['mode']}. {na.get('basis_note', '')}", 9, italic=True)
    elif t == "db_materials":
        db = numbering["db"]
        _table(doc, ["Id", "Element", "Specification", "Status", "Source"],
               [[m["id"], m["element"], m["spec"], m["status"], m.get("source", "")] for m in db["materials"]],
               widths=[28, 44, 48, 18, 32], caption="Materials (from design basis)")
    elif t == "calc":
        if b["ref"] not in results:
            raise BuildError(f"calc block references '{b['ref']}' but no results JSON with that id was supplied")
        render_calc(doc, results[b["ref"]])
    else:
        raise BuildError(f"unknown block type '{t}'")


# ----------------------------------------------------------------------------- main build
def build(content: dict, db: dict, results: dict, out_dir: pathlib.Path) -> dict:
    errors, warnings = check_content(content, db)
    if errors:
        raise BuildError("; ".join(errors))
    P, A = db["project"], db["approval"]
    d = content["doc"]
    style = "decimal" if d["type"] == "specification" else "lettered"
    draft = A["status"] != "APPROVED"

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(25), Mm(20)
    doc.styles["Normal"].font.name = FONT

    # header / footer
    hp = sec.header.paragraphs[0]
    _font(hp.add_run(f"{P['code']}  |  {d['number']}  |  Rev {P['revision']}  |  {P['issue_date']}"), 8, color=GREY)
    if draft:
        _font(hp.add_run("     DRAFT - NOT FOR ISSUE"), 8, bold=True, color=RGBColor(0xC0, 0, 0))
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("Page ")
    _font(r, 8, color=GREY)
    r2 = fp.add_run()
    _font(r2, 8, color=GREY)
    _field(r2, "PAGE")

    # title + control block
    _para(doc, d["title"].upper(), 16, bold=True, space_after=2)
    _para(doc, f"{P['code']}_{P['name']}", 11, space_after=10)
    _table(doc, ["Item", "Value"], [
        ["Document number", d["number"]], ["Project", f"{P['code']} {P['name']}"], ["Client", P["client"]],
        ["Stage", P["stage"]], ["Revision / date", f"{P['revision']} / {P['issue_date']}"],
        ["Design basis", f"{P['code']} design basis, status {A['status']}"],
        ["National Annex", f"{db['national_annex']['mode']}"
                           + (f" ({db['national_annex']['country']})" if db['national_annex'].get('country') else "")],
        ["Prepared by", A.get("prepared_by") or ""], ["Checked by", ""], ["Approved by", ""],
    ], widths=[45, 125], caption="Document control")
    if d.get("status_note"):
        _para(doc, d["status_note"], 9, italic=True)

    # body
    for si, s in enumerate(content["sections"]):
        num = str(si + 1) if style == "decimal" else _letters(si)
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.keep_with_next = True
        _font(h.add_run(f"{num}{'' if style == 'decimal' else '.'}  {s['title'].upper()}"), 11, bold=True)
        numbering = {"style": style, "section": num, "clause": 0, "db": db}
        for b in s["blocks"]:
            render_block(doc, b, results, numbering)

    # assumptions to be confirmed - always, from design basis
    si = len(content["sections"])
    num = str(si + 1) if style == "decimal" else _letters(si)
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    _font(h.add_run(f"{num}{'' if style == 'decimal' else '.'}  ASSUMPTIONS TO BE CONFIRMED"), 11, bold=True)
    items = [o for o in db.get("open_items", []) if o.get("status", "open") == "open"]
    _table(doc, ["Id", "Open item (from design basis)"], [[o["id"], o["description"]] for o in items] or [["-", "None recorded"]],
           widths=[15, 155])

    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{d['number']}_{P['revision']}"
    docx_path = out_dir / f"{stem}.docx"
    doc.save(docx_path)
    pdf_path = None
    if shutil.which("soffice"):
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                       check=True, capture_output=True, timeout=180)
        pdf_path = out_dir / f"{stem}.pdf"
        if not pdf_path.exists():
            raise BuildError("PDF conversion reported success but no PDF was produced")
    else:
        warnings.append("LibreOffice not found - PDF not produced")
    Document(docx_path)  # re-open: proves the DOCX is readable
    return {"docx": str(docx_path), "pdf": str(pdf_path) if pdf_path else None, "warnings": warnings}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("content")
    ap.add_argument("--design-basis", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--results", nargs="*", default=[])
    ns = ap.parse_args(argv)
    content = yaml.safe_load(open(ns.content, encoding="utf-8"))
    db = yaml.safe_load(open(ns.design_basis, encoding="utf-8"))
    results = {}
    for f in ns.results:
        r = json.load(open(f, encoding="utf-8"))
        results[r["id"]] = r
    try:
        res = build(content, db, results, pathlib.Path(ns.out))
    except BuildError as e:
        print(f"BUILD STOPPED: {e}", file=sys.stderr)
        return 1
    print(json.dumps(res, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
