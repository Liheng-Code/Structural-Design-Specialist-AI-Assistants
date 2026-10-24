"""Phase 4 tests: calc logic, results JSON, document generation. Run: python -m unittest discover -s tests -v"""
import copy
import json
import pathlib
import shutil
import subprocess
import sys
import tempfile
import unittest

import yaml
from docx import Document

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "04_tools"))
sys.path.insert(0, str(ROOT / "04_tools" / "docgen"))
from calc_lib.rc import ec2_anchorage as a  # noqa: E402
import build_document as bd  # noqa: E402

EX = ROOT / "03_projects" / "PJR-2026-003"
HAVE_SOFFICE = shutil.which("soffice") is not None


def db():
    return yaml.safe_load((EX / "design-basis.yaml").read_text(encoding="utf-8"))


class CalcLib(unittest.TestCase):
    """Hand values, worked independently of the code (C30/37, fyk 390)."""

    def test_fctd(self):
        self.assertAlmostEqual(a.fctd("C30/37"), 2.0 / 1.5, places=6)

    def test_bond_stress_good_and_poor(self):
        self.assertAlmostEqual(a.bond_stress("C30/37", True, 12).fbd, 3.0, places=6)   # 2.25*1*1*1.3333
        self.assertAlmostEqual(a.bond_stress("C30/37", False, 12).fbd, 2.1, places=6)  # x 0.7

    def test_eta2_large_bars(self):
        self.assertAlmostEqual(a.bond_stress("C30/37", True, 40).eta2, 0.92)

    def test_lb_rqd_hand_value(self):
        fyd = 390 / 1.15                                   # 339.1304 MPa
        self.assertAlmostEqual(a.fyd(390), fyd, places=4)
        self.assertAlmostEqual(a.lb_rqd(12, fyd, 3.0), 3 * fyd / 3.0, places=4)   # 339.13 mm
        self.assertAlmostEqual(a.lb_rqd(12, fyd, 3.0), 339.13, places=2)

    def test_rejects_untabulated_class_without_explicit_value(self):
        with self.assertRaises(KeyError):
            a.fctd("C90/105")
        self.assertAlmostEqual(a.fctd("C90/105", 3.0), 2.0)  # explicit value allowed

    def test_rejects_bad_input(self):
        with self.assertRaises(ValueError):
            a.bond_stress("C30/37", True, 0)
        with self.assertRaises(ValueError):
            a.lb_rqd(12, 339, 0)


class RunScript(unittest.TestCase):
    def run_check(self, inputs=None, extra=()):
        tmp = pathlib.Path(tempfile.mkdtemp())
        src = EX / "anchorage-inputs-sd004-r00-as-issued.yaml"
        if inputs is not None:
            src = tmp / "in.yaml"
            src.write_text(yaml.safe_dump(inputs), encoding="utf-8")
        out = tmp / "r.json"
        cp = subprocess.run([sys.executable, str(ROOT / "04_tools/calc_lib/runs/anchorage_check.py"), str(src), str(out), *extra],
                            capture_output=True, text=True)
        return cp, out

    def test_reconciles_with_drawing_on_drawing_basis(self):
        cp, out = self.run_check(extra=("--design-basis", str(EX / "design-basis.yaml")))
        self.assertEqual(cp.returncode, 0, cp.stderr)
        r = json.loads(out.read_text())
        self.assertTrue(r["reconciliation"]["all_within_tolerance_drawing_basis"])
        self.assertEqual(r["results"]["lb_rqd_good"][1], 339.1)

    def test_detects_mandrel_mismatch_in_sd004(self):
        cp, out = self.run_check()
        r = json.loads(out.read_text())
        self.assertEqual(r["reconciliation"]["mandrel"]["mismatch_bars"], [20, 25])

    def test_refuses_untraceable_design_basis_ref(self):
        inp = yaml.safe_load((EX / "anchorage-inputs-sd004-r00-as-issued.yaml").read_text(encoding="utf-8"))
        inp["design_basis_refs"] = ["M-DOES-NOT-EXIST"]
        cp, _ = self.run_check(inp, ("--design-basis", str(EX / "design-basis.yaml")))
        self.assertNotEqual(cp.returncode, 0)
        self.assertIn("untraceable", cp.stderr + cp.stdout)

    def test_wrong_drawing_value_is_flagged(self):
        inp = yaml.safe_load((EX / "anchorage-inputs-sd004-r00-as-issued.yaml").read_text(encoding="utf-8"))
        inp["compare_with"]["rows"]["good"][1] = 400   # deliberately wrong
        cp, out = self.run_check(inp)
        self.assertFalse(json.loads(out.read_text())["reconciliation"]["all_within_tolerance_drawing_basis"])


class DocGen(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = pathlib.Path(tempfile.mkdtemp())
        out = cls.tmp / "r.json"
        subprocess.run([sys.executable, str(ROOT / "04_tools/calc_lib/runs/anchorage_check.py"),
                        str(EX / "anchorage-inputs-sd004-r00-as-issued.yaml"), str(out)], check=True, capture_output=True)
        cls.results = {json.loads(out.read_text())["id"]: json.loads(out.read_text())}
        cls.spec = yaml.safe_load((EX / "content/specification.yaml").read_text(encoding="utf-8"))
        cls.rep = yaml.safe_load((EX / "content/calculation-report.yaml").read_text(encoding="utf-8"))

    def text(self, path):
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [c.text for r in t.rows for c in r.cells]
        for s in d.sections:
            parts += [p.text for p in s.header.paragraphs]
        return "\n".join(parts)

    def test_specification_builds_opens_and_is_stamped_draft(self):
        res = bd.build(self.spec, db(), {}, self.tmp / "spec")
        t = self.text(res["docx"])
        self.assertIn("DRAFT - NOT FOR ISSUE", t)
        self.assertIn("PJR-2026-003", t)             # project code comes from design basis
        self.assertIn("1.1", t)
        self.assertIn("4.1", t)                      # decimal numbering continuous to last section
        if HAVE_SOFFICE:
            self.assertTrue(pathlib.Path(res["pdf"]).stat().st_size > 5000)

    def test_report_contains_calc_and_open_items_and_blank_signoff(self):
        res = bd.build(self.rep, db(), self.results, self.tmp / "rep")
        t = self.text(res["docx"])
        self.assertIn("CALC-ANCH-001", t)
        self.assertIn("louvres", t)                  # F3 not truncated (regression)
        self.assertIn("site altitude", t)            # F4 not truncated (regression)
        self.assertIn("mandrel", t.lower())
        d = Document(res["docx"])
        ctrl = {r.cells[0].text: r.cells[1].text for r in d.tables[0].rows}
        self.assertEqual(ctrl["Checked by"], "")
        self.assertEqual(ctrl["Approved by"], "")

    def test_clause_without_source_stops_build(self):
        bad = copy.deepcopy(self.spec)
        bad["sections"][0]["blocks"].append({"type": "clause", "text": "Unsourced requirement."})
        with self.assertRaises(bd.BuildError):
            bd.build(bad, db(), {}, self.tmp / "bad")

    def test_missing_results_reference_stops_build(self):
        with self.assertRaises(bd.BuildError):
            bd.build(self.rep, db(), {}, self.tmp / "bad2")

    def test_approved_basis_removes_stamp_but_still_never_fills_signoff(self):
        d = db()
        d["approval"]["status"] = "APPROVED"
        res = bd.build(self.spec, d, {}, self.tmp / "appr")
        t = self.text(res["docx"])
        self.assertNotIn("DRAFT - NOT FOR ISSUE", t)
        ctrl = {r.cells[0].text: r.cells[1].text for r in Document(res["docx"]).tables[0].rows}
        self.assertEqual(ctrl["Approved by"], "")


if __name__ == "__main__":
    unittest.main()
