"""Phase 5 tests: Excel calculation workbook. Needs LibreOffice for recalculation (tests skip without it)."""
import json
import pathlib
import shutil
import sys
import tempfile
import unittest

import yaml
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools" / "xlsxgen"))
from calc_lib.rc import ec2_anchorage as a  # noqa: E402
import build_anchorage_workbook as bw  # noqa: E402
import verify_workbook as vw  # noqa: E402

EX = ROOT / "design-basis" / "examples" / "PRJ-2026-003"
HAVE = shutil.which("soffice") is not None


class HandValues(unittest.TestCase):
    """The expected values typed into the Checks sheet must agree with the independent python library."""

    def test_hand_values_match_library(self):
        fyd = a.fyd(390)
        lib = {"T1": a.fctd("C30/37"), "T2": fyd,
               "T3": a.bond_stress("C30/37", True, 12).fbd, "T4": a.bond_stress("C30/37", False, 12).fbd,
               "T5": a.lb_rqd(12, fyd, a.bond_stress("C30/37", True, 12).fbd),
               "T6": a.lb_rqd(12, fyd, a.bond_stress("C30/37", False, 12).fbd),
               "T7": 1.5 * a.lb_rqd(12, fyd, a.bond_stress("C30/37", True, 12).fbd),
               "T8": a.lb_rqd(25, fyd, a.bond_stress("C30/37", True, 25).fbd)}
        for tid, desc, exp, tol, ref, how in bw.HAND:
            self.assertAlmostEqual(lib[tid], exp, delta=max(tol, 1e-6), msg=tid)


@unittest.skipUnless(HAVE, "LibreOffice not installed")
class Workbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = pathlib.Path(tempfile.mkdtemp())
        cls.inp = yaml.safe_load((EX / "anchorage_inputs_SD004_R00_as_issued.yaml").read_text(encoding="utf-8"))
        cls.db = yaml.safe_load((EX / "design_basis.yaml").read_text(encoding="utf-8"))
        cls.res = json.loads((EX / "output" / "CALC-ANCH-001.json").read_text())
        cls.base = cls.tmp / "base.xlsx"
        bw.build(cls.inp, cls.db, cls.base)

    def fresh(self, name):
        p = self.tmp / name
        shutil.copyfile(self.base, p)
        return p

    def mutate(self, name, edits):
        p = self.fresh(name)
        wb = load_workbook(p)
        for sheet, cell, value in edits:
            wb[sheet][cell] = value
        wb.save(p)
        return p

    def test_baseline_verifies_and_ties_out_to_report(self):
        p = self.fresh("ok.xlsx")
        out = vw.verify(p, self.res)
        self.assertEqual(out["problems"], [])
        self.assertEqual(out["checks_passed"], 8)
        self.assertEqual(out["cover_status"], "CROSS-CHECKED - NOT ENGINEER-REVIEWED")

    def test_changing_input_changes_results_and_disables_baseline_checks(self):
        p = self.mutate("fyk.xlsx", [("Inputs", "C6", 500)])
        out = vw.verify(p)
        self.assertEqual(out["problems"], [])
        self.assertEqual(out["checks_passed"], 0)
        self.assertTrue(out["cover_status"].startswith("UNVERIFIED"))
        wv = load_workbook(p, data_only=True)
        fbd = a.bond_stress("C30/37", True, 12).fbd
        self.assertAlmostEqual(wv["Calc"]["E11"].value, a.lb_rqd(12, a.fyd(500), fbd), places=3)  # independent library

    def test_fctk_override_reproduces_drawing_basis(self):
        p = self.mutate("ovr.xlsx", [("Inputs", "C7", 2.03)])
        vw.verify(p)
        wv = load_workbook(p, data_only=True)
        self.assertAlmostEqual(wv["Calc"]["E11"].value, wv["Calc"]["H11"].value, places=6)

    def test_extra_bar_flows_through_without_errors(self):
        p = self.mutate("bar.xlsx", [("Inputs", "C35", 32)])
        out = vw.verify(p)
        self.assertEqual(out["problems"], [])
        wv = load_workbook(p, data_only=True)
        self.assertAlmostEqual(wv["Calc"]["B15"].value, 1.0)

    def test_large_bar_eta2(self):
        p = self.mutate("eta2.xlsx", [("Inputs", "C35", 40)])
        vw.verify(p)
        self.assertAlmostEqual(load_workbook(p, data_only=True)["Calc"]["B15"].value, 0.92, places=6)

    def test_typed_number_in_calc_is_detected(self):
        p = self.mutate("typed.xlsx", [("Calc", "E11", 339.1)])
        self.assertTrue(any("typed number" in x for x in vw.verify(p)["problems"]))

    def test_literal_in_calc_formula_is_detected(self):
        p = self.mutate("lit.xlsx", [("Calc", "G11", "=1.5*E11")])
        self.assertTrue(any("numeric literal" in x for x in vw.verify(p)["problems"]))

    def test_wrong_hand_value_makes_workbook_fail(self):
        p = self.mutate("bad.xlsx", [("Checks", "C9", 400.0)])
        out = vw.verify(p)
        self.assertTrue(any("FAILS" in x for x in out["problems"]))
        self.assertTrue(out["cover_status"].startswith("CHECKS FAILING"))

    def test_invalid_concrete_class_surfaces_error(self):
        p = self.mutate("cls.xlsx", [("Inputs", "C5", "C99/200")])
        self.assertTrue(any("#N/A" in x for x in vw.verify(p)["problems"]))

    def test_reviewer_name_changes_status_but_never_says_validated(self):
        p = self.mutate("rev.xlsx", [("Cover", "B12", "A. Engineer")])
        out = vw.verify(p)
        self.assertEqual(out["cover_status"], "REVIEWED BY A. Engineer")
        self.assertNotIn("VALIDATED", out["cover_status"].upper())

    def test_mandrel_error_in_drawing_is_flagged(self):
        p = self.fresh("mandrel.xlsx")
        vw.verify(p)
        wv = load_workbook(p, data_only=True)
        self.assertEqual([wv["Outputs"][f"D{r}"].value for r in range(19, 24)],
                         ["YES", "YES", "YES", "NO - CHECK", "NO - CHECK"])

    def test_inputs_have_validation_and_units(self):
        wb = load_workbook(self.base)
        self.assertTrue(wb["Inputs"].data_validations.dataValidation)
        for r in (5, 6, 10, 11, 12, 13):
            self.assertTrue(wb["Inputs"][f"E{r}"].value, f"source missing row {r}")


if __name__ == "__main__":
    unittest.main()
