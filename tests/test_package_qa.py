"""Phase 8 acceptance tests: cross-document QA on an integrated package and on deliberately broken packages.
Run: python -m unittest tests.test_package_qa -v. DXF cases need 01_standards/05-drafting-styles/05-04-layer-standards.csv."""
import csv
import json
import pathlib
import shutil
import subprocess
import sys
import tempfile
import unittest

import ezdxf
import yaml
from docx import Document
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "04_tools" / "qa"))
sys.path.insert(0, str(ROOT / "tests" / "fixtures"))
import package_qa as qa  # noqa: E402
import make_synthetic_etabs as fx  # noqa: E402

EX = ROOT / "03_projects" / "PJR-2026-003"
STD = ROOT / "01_standards"
HAVE_STD = (STD / "05-drafting-styles" / "05-04-layer-standards.csv").exists()
CR, SP, WB = "WH-STR-CR-001_R00", "WH-STR-SP-001_R00", "WH-STR-WB-001_R00"


def make_pkg():
    d = pathlib.Path(tempfile.mkdtemp()) / "pkg"
    d.mkdir()
    shutil.copytree(EX / "output", d / "output")
    for f in ("design-basis.yaml", "document-register.csv"):
        shutil.copy(EX / f, d / f)
    return d


def check(d, std=True):
    return qa.run(d, STD if (std and HAVE_STD) else None)


def hits(q, sev=None, check_id=None, text=None):
    return [f for f in q.f if (sev is None or f["severity"] == sev) and (check_id is None or f"[{check_id}]" in f["finding"])
            and (text is None or text in (f["finding"] + f["evidence"] + f["location"]))]


def gate(q):
    return qa.summarise(q)["consistency_gate_pass"]


def edit_docx(path, fn):
    d = Document(path)
    fn(d)
    d.save(path)


def set_cell(doc, label, value):
    for t in doc.tables:
        for r in t.rows:
            if r.cells[0].text == label:
                r.cells[1].text = value
                return
    raise KeyError(label)


def edit_db(d, fn):
    p = d / "design-basis.yaml"
    db = yaml.safe_load(p.read_text(encoding="utf-8"))
    fn(db)
    p.write_text(yaml.safe_dump(db, sort_keys=False, allow_unicode=True), encoding="utf-8")


class Integrated(unittest.TestCase):
    def test_example_package_passes_consistency_gate_but_not_issue_ready(self):
        q = check(make_pkg())
        self.assertTrue(gate(q), [f["finding"] for f in q.f if f["severity"] in ("Critical", "Major")])
        self.assertEqual(hits(q, "Critical"), [])
        s = qa.summarise(q)
        self.assertTrue(s["not_checked"], "limits must be listed")

    def test_report_files_written_and_never_approve(self):
        d = make_pkg()
        q = check(d)
        s = qa.write(q, d / "qa", "pkg", "DRAFT")
        with open(d / "qa" / "issue-register.csv", encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        self.assertEqual(list(rows[0]), qa.COLS)
        self.assertTrue(all(r["review_status"] == "Open" for r in rows))
        txt = (d / "qa" / "qa_report.md").read_text(encoding="utf-8")
        self.assertIn("NOT READY", txt)
        self.assertIn("authorised engineering review", txt)
        self.assertNotIn("Accepted-by-engineer", " ".join(r["review_status"] for r in rows))

    def test_cli_exit_codes(self):
        d = make_pkg()
        cmd = [sys.executable, str(ROOT / "04_tools/qa/package_qa.py"), str(d), "--out", str(d / "qa")]
        self.assertEqual(subprocess.run(cmd, capture_output=True).returncode, 0)
        (d / "design-basis.yaml").unlink()
        self.assertEqual(subprocess.run(cmd, capture_output=True).returncode, 1)


class Incomplete(unittest.TestCase):
    def test_missing_design_basis(self):
        d = make_pkg()
        (d / "design-basis.yaml").unlink()
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Critical", "Q1"))

    def test_missing_register(self):
        d = make_pkg()
        (d / "document-register.csv").unlink()
        self.assertTrue(hits(check(d), "Major", "Q2"))

    def test_empty_output(self):
        d = make_pkg()
        shutil.rmtree(d / "output")
        q = check(d)
        self.assertFalse(gate(q))

    def test_results_file_missing(self):
        d = make_pkg()
        (d / "output" / "CALC-ANCH-001.json").unlink()
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Major", "Q7"))
        self.assertTrue(any("Q7" in n for n in q.not_checked))

    def test_no_verification_record(self):
        d = make_pkg()
        p = d / "output" / "CALC-ANCH-001.json"
        j = json.loads(p.read_text())
        j.pop("verification")
        p.write_text(json.dumps(j))
        self.assertTrue(hits(check(d), "Critical", "Q7", "verification"))

    def test_corrupt_docx(self):
        d = make_pkg()
        (d / "output" / f"{CR}.docx").write_bytes(b"not a docx")
        self.assertTrue(hits(check(d), "Critical", "Q11"))

    def test_sheets_without_files_are_declared_not_silently_passed(self):
        q = check(make_pkg())
        self.assertTrue(any("register sheets have no file" in n for n in q.not_checked))


class Inconsistent(unittest.TestCase):
    def test_report_number_differs_from_results(self):
        d = make_pkg()

        def f(doc):
            for t in doc.tables:
                for r in t.rows:
                    if r.cells[0].text == "lb,rqd good bond (mm)":
                        r.cells[2].text = "350"
        edit_docx(d / "output" / f"{CR}.docx", f)
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Critical", "Q7", "lb,rqd good"))

    def test_results_changed_after_report(self):
        d = make_pkg()
        p = d / "output" / "CALC-ANCH-001.json"
        j = json.loads(p.read_text())
        j["results"]["lap_basic_good"][4] += 40
        p.write_text(json.dumps(j))
        q = check(d)
        self.assertTrue(hits(q, "Critical", "Q7", "lap l0"))
        self.assertTrue(hits(q, "Critical", "Q10"), "workbook tie-out must also fail")

    def test_workbook_cached_value_differs(self):
        d = make_pkg()
        p = d / "output" / f"{WB}.xlsx"
        wb = load_workbook(p)
        wb["Calc"]["E10"] = 999     # typed number in Calc
        wb.save(p)
        self.assertTrue(hits(check(d), "Critical", "Q10"))

    def test_client_differs(self):
        d = make_pkg()
        edit_docx(d / "output" / f"{SP}.docx", lambda doc: set_cell(doc, "Client", "Mr. YYY"))
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Major", "Q4", "Client"))

    def test_revision_differs(self):
        d = make_pkg()
        edit_docx(d / "output" / f"{CR}.docx", lambda doc: set_cell(doc, "Revision / date", "R01 / 13-OCT-2026"))
        self.assertTrue(hits(check(d), "Major", "Q4", "Revision"))

    def test_workbook_cover_stage_differs(self):
        d = make_pkg()
        p = d / "output" / f"{WB}.xlsx"
        wb = load_workbook(p)
        for r in wb["Cover"].iter_rows():
            if r[0].value == "Stage":
                r[1].value = "FOR CONSTRUCTION"
        wb.save(p)
        self.assertTrue(hits(check(d), "Major", "Q4", "Stage"))

    def test_material_differs_from_design_basis(self):
        d = make_pkg()

        def f(db):
            for m in db["materials"]:
                if m["id"] == "M-REBAR":
                    m["spec"] = "fyk = 500 MPa deformed bars"
        edit_db(d, f)
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Major", "Q8", "fyk"))

    def test_standard_not_in_register(self):
        d = make_pkg()
        p = d / "document-register.csv"
        with open(p, encoding="utf-8") as fh:
            rows = [r for r in csv.DictReader(fh) if not r["sheet_number"].endswith("-001")]
        with open(p, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=rows[0].keys())
            w.writeheader()
            w.writerows(rows)
        self.assertTrue(hits(check(d), "Major", "Q3", "register"))

    def test_register_revision_differs_from_file(self):
        d = make_pkg()
        p = d / "document-register.csv"
        with open(p, encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        for r in rows:
            if r["sheet_number"].endswith("-002"):
                r["revision"] = "R01"
        with open(p, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=rows[0].keys())
            w.writeheader()
            w.writerows(rows)
        self.assertFalse(gate(check(d)))

    def test_project_code_differs(self):
        d = make_pkg()
        edit_db(d, lambda db: db["project"].__setitem__("code", "PJR-2026-004"))
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Major", "Q4"))

    @unittest.skipUnless(HAVE_STD, "company standards not found")
    def test_dxf_title_block_revision_differs(self):
        d = make_pkg()
        p = d / "output" / "cad" / "WH-STR-SD-001_R00.dxf"
        doc = ezdxf.readfile(p)
        for e in doc.modelspace().query("INSERT"):
            for a in e.attribs:
                if a.dxf.tag == "REVISION":
                    a.dxf.text = "R05"
        doc.saveas(p)
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Major", "Q9") or hits(q, "Major", "Q4"))

    @unittest.skipUnless(HAVE_STD, "company standards not found")
    def test_dxf_without_standards_is_declared_unchecked(self):
        q = check(make_pkg(), std=False)
        self.assertTrue(any(n.startswith("Q9") for n in q.not_checked))


class UnsupportedClaims(unittest.TestCase):
    def test_workbook_claims_validated(self):
        d = make_pkg()
        p = d / "output" / f"{WB}.xlsx"
        wb = load_workbook(p)
        for r in wb["Cover"].iter_rows():
            if r[0].value == "STATUS":
                r[1].value = "VALIDATED"
        wb.save(p)
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Critical", "Q5"))

    def test_document_claims_approval(self):
        d = make_pkg()
        edit_docx(d / "output" / f"{SP}.docx", lambda doc: doc.add_paragraph("This specification is APPROVED FOR ISSUE."))
        self.assertTrue(hits(check(d), "Critical", "Q5"))

    def test_draft_stamp_removed(self):
        d = make_pkg()

        def f(doc):
            for s in doc.sections:
                for p in s.header.paragraphs:
                    for r in p.runs:
                        r.text = r.text.replace("DRAFT - NOT FOR ISSUE", "")
        edit_docx(d / "output" / f"{CR}.docx", f)
        self.assertTrue(hits(check(d), "Major", "Q5"))

    def test_new_placeholder_in_document(self):
        d = make_pkg()
        edit_docx(d / "output" / f"{CR}.docx", lambda doc: doc.add_paragraph("Soil bearing capacity TBC."))
        self.assertTrue(hits(check(d), "Major", "Q6", "TBC"))

    def test_inherited_placeholder_is_info_in_draft(self):
        q = check(make_pkg())
        self.assertTrue(hits(q, "Info", "Q6", "XXX"))
        self.assertTrue(hits(q, "Minor", "Q6", "client"))

    def test_placeholder_in_non_draft_basis_is_major(self):
        d = make_pkg()
        edit_db(d, lambda db: db["approval"].__setitem__("status", "FOR_REVIEW"))
        self.assertTrue(hits(check(d), "Major", "Q6", "client"))

    def test_approved_basis_with_open_items_and_blank_editions(self):
        d = make_pkg()
        edit_db(d, lambda db: db["approval"].update(status="APPROVED"))
        q = check(d)
        self.assertTrue(hits(q, "Major", "Q13"))

    def test_filled_checked_by_is_flagged_for_verification(self):
        d = make_pkg()
        edit_docx(d / "output" / f"{CR}.docx", lambda doc: set_cell(doc, "Checked by", "A. Engineer"))
        self.assertTrue(hits(q := check(d), "Info", "Q13", "Checked by"))
        self.assertEqual(hits(q, "Critical"), [])


class AnalysisImport(unittest.TestCase):
    def pkg_with_import(self, **defects):
        sys.path.insert(0, str(ROOT / "04_tools" / "etabs"))
        import etabs_import as ei
        d = make_pkg()
        src = pathlib.Path(tempfile.mkdtemp())
        fx.build(src, **defects)
        ei.run([src / "synthetic_export.xlsx"], d / "analysis" / "SYN-1", "SYN-1", members_csv=src / "members.csv",
               expected_yaml=src / "expected_cases.yaml", hand_yaml=src / "hand_totals.yaml", software="synthetic", model_file="synthetic.edb", run_date="14-OCT-2026")
        return d

    def test_good_import_without_design_basis_provenance_is_major(self):
        d = self.pkg_with_import()
        q = check(d)
        self.assertTrue(hits(q, "Major", "Q12", "analysis_provenance"))
        self.assertFalse(hits(q, "Critical", "Q12"))

    def test_provenance_recorded_passes_q12(self):
        d = self.pkg_with_import()
        lg = json.loads((d / "analysis/SYN-1/import_log.json").read_text())
        prov = {"software": "synthetic", "model_file": "synthetic.edb", "export_sha256": {e["file"]: e["sha256"] for e in lg["provenance"]["export_files"]}}
        edit_db(d, lambda db: db.__setitem__("analysis_provenance", prov))
        q = check(d)
        self.assertEqual(hits(q, check_id="Q12"), [])

    def test_failed_import_is_critical(self):
        d = self.pkg_with_import(scale_base=True)
        q = check(d)
        self.assertFalse(gate(q))
        self.assertTrue(hits(q, "Critical", "Q12"))

    def test_tampered_raw_export(self):
        d = self.pkg_with_import()
        with open(d / "analysis/SYN-1/raw/synthetic_export.xlsx", "ab") as fh:
            fh.write(b"x")
        self.assertTrue(hits(check(d), "Critical", "Q12", "SHA-256"))

    def test_db_model_definition_import_is_major_not_critical(self):
        sys.path.insert(0, str(ROOT / "04_tools" / "etabs"))
        import etabs_db_import as edb
        import make_synthetic_etabs_db as fdb
        d = make_pkg()
        src = pathlib.Path(tempfile.mkdtemp())
        fdb.build(str(src / "m.db"))
        edb.run(src / "m.db", d / "analysis" / "SYN-DB", "SYN-DB")
        q = check(d)
        self.assertTrue(hits(q, "Major", "Q12", "model definition only"))
        self.assertEqual(hits(q, "Critical", "Q12"), [])
        with open(d / "analysis/SYN-DB/raw/m.db", "ab") as fh:
            fh.write(b"x")
        self.assertTrue(hits(check(d), "Critical", "Q12", "SHA-256"))

    def test_unperformed_import_check_is_major(self):
        sys.path.insert(0, str(ROOT / "04_tools" / "etabs"))
        import etabs_import as ei
        d = make_pkg()
        src = pathlib.Path(tempfile.mkdtemp())
        fx.build(src)
        ei.run([src / "synthetic_export.xlsx"], d / "analysis" / "SYN-1", "SYN-1")   # no expected cases / members / hand totals
        q = check(d)
        self.assertTrue(hits(q, "Major", "Q12", "not performed"))


if __name__ == "__main__":
    unittest.main()
